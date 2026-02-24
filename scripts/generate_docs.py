#!/usr/bin/env python3
"""
Script to convert the research Markdown document to DOCX and PDF formats.

Usage:
    python scripts/generate_docs.py

Generates:
    docs/research/consumo_alcohol_espana.docx
    docs/research/consumo_alcohol_espana.pdf
"""

import os
import re
import markdown
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)

NUMBERED_LIST_RE = re.compile(r"^(\d+)\.\s+(.+)$")
RESEARCH_DIR = os.path.join(PROJECT_ROOT, "docs", "research")
MD_FILE = os.path.join(RESEARCH_DIR, "consumo_alcohol_espana.md")
DOCX_FILE = os.path.join(RESEARCH_DIR, "consumo_alcohol_espana.docx")
PDF_FILE = os.path.join(RESEARCH_DIR, "consumo_alcohol_espana.pdf")


def read_markdown():
    with open(MD_FILE, "r", encoding="utf-8") as f:
        return f.read()


def parse_table(lines):
    """Parse a Markdown table into a list of rows (each row is a list of cells)."""
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        if not cells or all(re.match(r"^[-:]+$", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def clean_md(text):
    """Remove markdown formatting from text."""
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = re.sub(r"~", "", text)
    return text


def add_rich_paragraph(doc, text, style=None):
    """Add a paragraph with bold parts preserved from Markdown."""
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    p = doc.add_paragraph(style=style)
    bold_pattern = re.compile(r"\*\*([^*]+)\*\*")
    last_end = 0
    for m in bold_pattern.finditer(text):
        if m.start() > last_end:
            p.add_run(text[last_end:m.start()])
        run = p.add_run(m.group(1))
        run.bold = True
        last_end = m.end()
    if last_end < len(text):
        p.add_run(text[last_end:])
    return p


def generate_docx(md_content):
    """Generate a DOCX file from the Markdown content."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        # Headings
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = clean_md(stripped.lstrip("#").strip())
            if level == 1:
                heading = doc.add_heading(text, level=0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # Tables
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table(table_lines)
            if rows:
                num_cols = len(rows[0])
                table = doc.add_table(rows=len(rows), cols=num_cols)
                table.style = "Light Grid Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for row_idx, row_data in enumerate(rows):
                    for col_idx, cell_text in enumerate(row_data):
                        if col_idx < num_cols:
                            cell = table.cell(row_idx, col_idx)
                            cell.text = clean_md(cell_text)
                            if row_idx == 0:
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.bold = True
                doc.add_paragraph()
            continue

        # Bullet points
        if stripped.startswith("- ") or stripped.startswith("* "):
            text = stripped[2:]
            indent_level = len(line) - len(line.lstrip())
            bullet_style = "List Bullet 2" if indent_level > 0 else "List Bullet"
            add_rich_paragraph(doc, text, style=bullet_style)
            i += 1
            continue

        # Numbered lists
        match = NUMBERED_LIST_RE.match(stripped)
        if match:
            add_rich_paragraph(doc, match.group(2), style="List Number")
            i += 1
            continue

        # Italic footer text
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            text = stripped.strip("*")
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.italic = True
            run.font.size = Pt(9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Regular paragraphs
        add_rich_paragraph(doc, stripped)
        i += 1

    doc.save(DOCX_FILE)
    print(f"DOCX generated: {DOCX_FILE}")


def generate_pdf(md_content):
    """Generate a PDF file from the Markdown content using weasyprint."""
    from weasyprint import HTML

    html_content = markdown.markdown(
        md_content,
        extensions=["tables", "toc", "fenced_code"],
    )

    full_html = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <style>
        @page {{
            size: A4;
            margin: 2cm;
        }}
        body {{
            font-family: Calibri, Arial, sans-serif;
            font-size: 11pt;
            line-height: 1.5;
            color: #333;
        }}
        h1 {{
            color: #1a3c6e;
            text-align: center;
            font-size: 22pt;
            border-bottom: 2px solid #1a3c6e;
            padding-bottom: 10px;
            margin-bottom: 20px;
        }}
        h2 {{
            color: #1a3c6e;
            font-size: 16pt;
            border-bottom: 1px solid #ccc;
            padding-bottom: 5px;
            margin-top: 30px;
        }}
        h3 {{
            color: #2c5f9e;
            font-size: 13pt;
            margin-top: 20px;
        }}
        h4 {{
            color: #3a7bc8;
            font-size: 12pt;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 15px 0;
            font-size: 10pt;
        }}
        th {{
            background-color: #1a3c6e;
            color: white;
            padding: 8px 12px;
            text-align: left;
            border: 1px solid #1a3c6e;
        }}
        td {{
            padding: 6px 12px;
            border: 1px solid #ddd;
        }}
        tr:nth-child(even) {{
            background-color: #f5f7fa;
        }}
        ul, ol {{
            margin: 10px 0;
        }}
        li {{
            margin: 4px 0;
        }}
        strong {{
            color: #1a3c6e;
        }}
        hr {{
            border: none;
            border-top: 1px solid #ccc;
            margin: 30px 0;
        }}
        em {{
            font-size: 9pt;
            color: #666;
        }}
        a {{
            color: #2c5f9e;
        }}
    </style>
</head>
<body>
{html_content}
</body>
</html>"""

    HTML(string=full_html).write_pdf(PDF_FILE)
    print(f"PDF generated: {PDF_FILE}")


if __name__ == "__main__":
    md_content = read_markdown()
    generate_docx(md_content)
    generate_pdf(md_content)
    print("\nDone! Both files generated successfully.")
